"""
Verificación reproducible de la Parte 2 sobre los datos reales de RetailPro.
Lee Fact_Ventas (hoja 'ventas', columna total_venta) del dataset del proyecto
y recalcula media, mediana e IQR con el mismo criterio que PERCENTILE_CONT en SQL.

Uso:  python verificar_stats.py
"""
import math, statistics
import openpyxl

XLSX = "../pipeline-etl-techstore-powerbi/data/Pipeline_ETL_Dataset.xlsx"

def percentile_cont(sorted_vals, p):
    """Interpolación lineal — equivalente a PERCENTILE_CONT (SQL) y numpy 'linear'."""
    k = p * (len(sorted_vals) - 1)
    f, c = math.floor(k), math.ceil(k)
    if f == c:
        return sorted_vals[int(k)]
    return sorted_vals[f] + (k - f) * (sorted_vals[c] - sorted_vals[f])

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["ventas"]
headers = [c.value for c in ws[1]]
i = headers.index("total_venta")
vals = [float(r[i]) for r in ws.iter_rows(min_row=2, values_only=True) if r[i] is not None]
s = sorted(vals)

media = statistics.mean(vals)
mediana = percentile_cont(s, 0.5)
q1, q3 = percentile_cont(s, 0.25), percentile_cont(s, 0.75)
iqr = q3 - q1
li, ls = q1 - 1.5 * iqr, q3 + 1.5 * iqr
outliers = [v for v in s if v < li or v > ls]

print(f"n              = {len(vals)}")
print(f"media          = {media:.2f}")
print(f"mediana        = {mediana:.2f}")
print(f"Q1             = {q1:.2f}")
print(f"Q3             = {q3:.2f}")
print(f"IQR            = {iqr:.2f}")
print(f"limite inf/sup = {li:.2f} / {ls:.2f}")
print(f"outliers ({len(outliers)})   = {[round(v, 2) for v in outliers]}")
