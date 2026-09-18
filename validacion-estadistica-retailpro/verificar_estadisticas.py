"""
Verifica los cálculos de la Parte 2 (validación estadística de RetailPro)
sobre la columna total_venta de la tabla ventas del dataset real del proyecto.

Reproduce en Python lo que en SQL se resuelve con AVG, PERCENTILE_CONT(0.25/0.5/0.75)
y el método IQR. PERCENTILE_CONT usa interpolación lineal (rango = p*(n-1)),
igual que PostgreSQL.

Uso:  python3 verificar_estadisticas.py
Requiere:  pip install openpyxl
"""
import statistics
from pathlib import Path
import openpyxl

XLSX = Path(__file__).resolve().parents[0] / ".." / \
    "pipeline-etl-techstore-powerbi" / "data" / "Pipeline_ETL_Dataset.xlsx"


def percentile_cont(data, p):
    """PERCENTILE_CONT: interpolación lineal, rango = p*(n-1) (0-indexed)."""
    s = sorted(data)
    n = len(s)
    if n == 1:
        return s[0]
    rank = p * (n - 1)
    lo = int(rank)
    frac = rank - lo
    if lo + 1 < n:
        return s[lo] + frac * (s[lo + 1] - s[lo])
    return s[lo]


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["ventas"]
    headers = [c.value for c in ws[1]]
    idx = headers.index("total_venta")
    vals = [float(r[idx]) for r in ws.iter_rows(min_row=2, values_only=True)
            if r[idx] is not None]

    n = len(vals)
    media = sum(vals) / n
    mediana = percentile_cont(vals, 0.5)
    q1 = percentile_cont(vals, 0.25)
    q3 = percentile_cont(vals, 0.75)
    iqr = q3 - q1
    lim_inf = q1 - 1.5 * iqr
    lim_sup = q3 + 1.5 * iqr
    outliers = sorted(v for v in vals if v < lim_inf or v > lim_sup)

    print(f"n                = {n}")
    print(f"media (AVG)      = {media:.2f}")
    print(f"mediana (P50)    = {mediana:.2f}")
    print(f"desv. muestral   = {statistics.stdev(vals):.2f}")
    print(f"Q1               = {q1:.2f}")
    print(f"Q3               = {q3:.2f}")
    print(f"IQR              = {iqr:.2f}")
    print(f"limite inferior  = {lim_inf:.2f}")
    print(f"limite superior  = {lim_sup:.2f}")
    print(f"outliers ({len(outliers)})     = {outliers}")
    print(f"ratio media/mediana = {media / mediana:.2f}")
    pct = 100 * sum(outliers) / sum(vals)
    print(f"outliers = {len(outliers)/n*100:.0f}% de las ventas, "
          f"{pct:.0f}% de la facturacion")


if __name__ == "__main__":
    main()
