import openpyxl, sqlite3, datetime, json
XLSX="../pipeline-etl-techstore-powerbi/data/Pipeline_ETL_Dataset.xlsx"
wb=openpyxl.load_workbook(XLSX, data_only=True)
def sheet(name):
    ws=wb[name]; hdr=[c.value for c in ws[1]]
    return hdr,[list(r) for r in ws.iter_rows(min_row=2,values_only=True)]
ch,clientes=sheet("clientes"); ph,productos=sheet("productos"); vh,ventas=sheet("ventas"); kh,cats=sheet("categorias")
def dedupe(rows,i):
    s=set();o=[]
    for r in rows:
        if r[i] in s: continue
        s.add(r[i]);o.append(r)
    return o
clientes=dedupe(clientes,ch.index("id_cliente")); productos=dedupe(productos,ph.index("id_producto"))
con=sqlite3.connect(":memory:");cur=con.cursor()
def mk(n,h,rows):
    cur.execute(f'CREATE TABLE {n} ({",".join(chr(34)+c+chr(34) for c in h)})')
    cur.executemany(f'INSERT INTO {n} VALUES ({("?,"*len(h))[:-1]})',[[(v.isoformat() if isinstance(v,(datetime.date,datetime.datetime)) else v) for v in r] for r in rows])
mk("clientes",ch,clientes);mk("productos",ph,productos);mk("ventas",vh,ventas);mk("categorias",kh,cats)
con.commit()
def q(s): return cur.execute(s).fetchall()

out={}
out["total_ventas"]=q("SELECT ROUND(SUM(total_venta),2) FROM ventas")[0][0]
out["n_ventas"]=q("SELECT COUNT(*) FROM ventas")[0][0]
out["n_clientes"]=q("SELECT COUNT(*) FROM clientes")[0][0]
out["n_productos"]=q("SELECT COUNT(*) FROM productos")[0][0]
out["ticket_prom"]=round(q("SELECT AVG(total_venta) FROM ventas")[0][0],2)
# mediana
vals=sorted(r[0] for r in q("SELECT total_venta FROM ventas"))
import statistics
out["ticket_med"]=round(statistics.median(vals),2)
out["ventas_online"]=q("SELECT ROUND(SUM(total_venta),2) FROM ventas WHERE canal='Online'")[0][0]
# por mes
out["por_mes"]=q("SELECT substr(fecha_venta,1,7) m, ROUND(SUM(total_venta),2) FROM ventas GROUP BY m ORDER BY m")
# por año
out["por_anio"]=q("SELECT substr(fecha_venta,1,4) y, ROUND(SUM(total_venta),2) FROM ventas GROUP BY y ORDER BY y")
# por categoria (join productos)
out["por_cat"]=q("""SELECT COALESCE(p.categoria,'Sin Categoria') cat, ROUND(SUM(v.total_venta),2)
    FROM ventas v LEFT JOIN productos p ON p.id_producto=v.id_producto GROUP BY cat ORDER BY 2 DESC""")
# por canal
out["por_canal"]=q("SELECT canal, ROUND(SUM(total_venta),2), COUNT(*) FROM ventas GROUP BY canal ORDER BY 2 DESC")
# por pais (join clientes)
out["por_pais"]=q("""SELECT c.pais, ROUND(SUM(v.total_venta),2) FROM ventas v
    JOIN clientes c ON c.id_cliente=v.id_cliente GROUP BY c.pais ORDER BY 2 DESC""")
# top clientes
out["top_cli"]=q("""SELECT c.nombre_cliente, ROUND(SUM(v.total_venta),2) FROM ventas v
    JOIN clientes c ON c.id_cliente=v.id_cliente GROUP BY c.nombre_cliente ORDER BY 2 DESC LIMIT 6""")
# top productos
out["top_prod"]=q("""SELECT p.nombre_producto, ROUND(SUM(v.total_venta),2) FROM ventas v
    JOIN productos p ON p.id_producto=v.id_producto GROUP BY p.nombre_producto ORDER BY 2 DESC LIMIT 6""")
out["dist"]=vals
# crecimiento anual
pa=dict(out["por_anio"])
if '2023' in pa and '2024' in pa:
    out["crec"]=round((pa['2024']-pa['2023'])/pa['2023']*100,1)
print(json.dumps(out,indent=1,ensure_ascii=False,default=str))
json.dump(out,open("agg.json","w"),default=str)
