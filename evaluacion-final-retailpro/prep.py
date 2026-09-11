import openpyxl, sqlite3, datetime
XLSX="../pipeline-etl-techstore-powerbi/data/Pipeline_ETL_Dataset.xlsx"
wb=openpyxl.load_workbook(XLSX, data_only=True)
def sheet(name):
    ws=wb[name]; hdr=[c.value for c in ws[1]]
    return hdr,[list(r) for r in ws.iter_rows(min_row=2,values_only=True)]

ch,clientes=sheet("clientes")
ph,productos=sheet("productos")
vh,ventas=sheet("ventas")
kh,cats=sheet("categorias")
print("raw counts:",len(clientes),len(productos),len(ventas),len(cats))

# CLEAN: dedupe clientes/productos by id (keep first)
def dedupe(rows, ididx):
    seen=set(); out=[]
    for r in rows:
        if r[ididx] in seen: continue
        seen.add(r[ididx]); out.append(r)
    return out
clientes=dedupe(clientes, ch.index("id_cliente"))
productos=dedupe(productos, ph.index("id_producto"))
print("clean counts:",len(clientes),len(productos),len(ventas),len(cats))

con=sqlite3.connect(":memory:")
cur=con.cursor()
def mk(name,hdr,rows):
    cols=",".join(f'"{c}"' for c in hdr)
    cur.execute(f'CREATE TABLE {name} ({cols})')
    q="?,"*len(hdr)
    cur.executemany(f'INSERT INTO {name} VALUES ({q[:-1]})',[[ (v.isoformat() if isinstance(v,(datetime.date,datetime.datetime)) else v) for v in r] for r in rows])
mk("clientes",ch,clientes); mk("productos",ph,productos); mk("ventas",vh,ventas); mk("categorias",kh,cats)
con.commit()

# max fecha
cur.execute("SELECT MAX(date(fecha_venta)) FROM ventas"); maxf=cur.fetchone()[0]
cur.execute("SELECT MIN(date(fecha_venta)) FROM ventas"); minf=cur.fetchone()[0]
print("rango fechas:",minf,"->",maxf)

# Ejercicio 1: ultimos 30 dias desde la ultima venta
q1="""
SELECT c.nombre_cliente AS cliente,
       date(v.fecha_venta) AS fecha_compra,
       v.total_venta AS total
FROM ventas v
JOIN clientes c ON c.id_cliente = v.id_cliente
WHERE date(v.fecha_venta) >= date((SELECT MAX(fecha_venta) FROM ventas), '-30 day')
ORDER BY v.fecha_venta DESC;
"""
rows=cur.execute(q1).fetchall()
print("\nEJ1 filas ultimos 30 dias:",len(rows))
for r in rows: print(r)
