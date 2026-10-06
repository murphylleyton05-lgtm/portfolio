"""
Arma el Excel para armar el dashboard a mano en cualquier version de Tableau.

Trae precalculado lo que en el libro generado eran campos calculados
(grupo de la dona, marca de ultimo mes), asi en Tableau casi no hay formulas.
"""

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

RAIZ = Path(__file__).resolve().parents[1]
CSV = RAIZ / "data" / "vaca_muerta_operadores_mensual.csv"
XLSX = RAIZ / "data" / "VacaMuerta_Tableau.xlsx"
GRUPOS = {"YPF": "YPF", "VISTA ENERGY": "Vista", "PLUSPETROL": "Pluspetrol",
          "SHELL": "Shell", "PAN AMERICAN ENERGY": "PAE"}


def main() -> None:
    with CSV.open(encoding="utf-8") as f:
        filas = list(csv.DictReader(f))
    ultimo = max(f["Mes"] for f in filas)

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    cab = ["Mes", "Operador", "Grupo_Dona", "Petroleo_bbl_d", "Petroleo_anio_anterior_bbl_d",
           "Gas_Mm3_d", "Pozos_activos", "Ultimo_mes"]
    ws.append(cab)
    for f in filas:
        a, m, d = map(int, f["Mes"].split("-"))
        ant = f["Petroleo_bbl_d_anio_anterior"]
        ws.append([
            date(a, m, d), f["Operador"], GRUPOS.get(f["Operador"], "Resto"),
            int(f["Petroleo_bbl_d"]), int(ant) if ant else None,
            float(f["Gas_Mm3_d"]), int(f["Pozos_activos"]),
            "Si" if f["Mes"] == ultimo else "No",
        ])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for fila in ws.iter_rows(min_row=2, max_col=1):
        fila[0].number_format = "yyyy-mm-dd"
    for col, ancho in zip("ABCDEFGH", [12, 30, 12, 16, 28, 12, 14, 11]):
        ws.column_dimensions[col].width = ancho
    wb.save(XLSX)
    print(f"{len(filas)} filas -> {XLSX.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
